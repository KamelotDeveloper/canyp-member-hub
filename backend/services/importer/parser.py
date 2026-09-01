"""File parsing for the bulk import engine.

Supports CSV (delimiter auto-detected among comma/semicolon/tab) and XLSX
(first worksheet only via openpyxl). CSV text is decoded with an encoding
cascade: UTF-8-sig -> UTF-8 -> Latin-1 (first that succeeds without error).
"""

from __future__ import annotations

import csv
import io
from typing import Any

from backend.schemas.import_bulk import MAX_PARSEABLE_ROWS

SUPPORTED_EXTENSIONS = {".csv", ".xlsx"}
ROW_CAP = MAX_PARSEABLE_ROWS  # 10,000 parseable rows


class ParseError(Exception):
    """Raised when a file cannot be parsed or exceeds engine limits."""


def _detect_delimiter(sample: str) -> str:
    """Auto-detect CSV delimiter among comma, semicolon, and tab.

    Chooses the candidate with the most occurrences in the sample so a
    semicolon-delimited file whose values contain commas is not misrouted
    to comma. Ties resolve in standard order (comma, semicolon, tab).
    """
    counts = {",": sample.count(","), ";": sample.count(";"), "\t": sample.count("\t")}
    if not any(counts.values()):
        return ","
    best = ","
    best_count = -1
    for delim in (",", ";", "\t"):
        if counts[delim] > best_count:
            best = delim
            best_count = counts[delim]
    return best


def _decode_utf8_sig(raw: bytes) -> str | None:
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return None


def _decode_utf8(raw: bytes) -> str | None:
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError:
        return None


def _decode_latin1(raw: bytes) -> str:
    # Latin-1 never fails to decode any byte sequence.
    return raw.decode("latin-1")


def decode_bytes(raw: bytes) -> str:
    """Decode raw bytes using the encoding cascade UTF-8-sig -> UTF-8 -> Latin-1."""
    for decoder in (_decode_utf8_sig, _decode_utf8):
        text = decoder(raw)
        if text is not None:
            return text
    return _decode_latin1(raw)


def _parse_csv(text: str) -> list[list[str]]:
    """Parse CSV text into rows of raw string cells (header row first)."""
    sample = text[:2000]
    delimiter = _detect_delimiter(sample)
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows: list[list[str]] = []
    for row in reader:
        if not row or all(cell.strip() == "" for cell in row):
            continue
        rows.append([cell for cell in row])
    return rows


def _parse_xlsx(raw: bytes) -> list[list[str]]:
    """Parse the first worksheet of an XLSX workbook into raw string rows."""
    import io as _io

    from openpyxl import load_workbook

    try:
        wb = load_workbook(filename=_io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several exception types on bad input
        raise ParseError("Could not read the XLSX file.") from exc

    ws = wb.worksheets[0]
    rows: list[list[str]] = []
    for row in ws.iter_rows(values_only=True):
        cells = ["" if cell is None else str(cell).strip() for cell in row]
        if not any(cells):
            continue
        rows.append(cells)
    wb.close()
    return rows


def parse_file(filename: str, raw: bytes) -> list[list[str]]:
    """Parse an uploaded file into raw rows of string cells (header first).

    Returns a list where each element is a row and the first element is the
    header row. Raises ``ParseError`` on unsupported extension, unreadable
    content, or exceeding the parseable-row cap.
    """
    name = (filename or "").lower()
    if not name:
        raise ParseError("No file name provided.")
    ext = "." + name.rsplit(".", 1)[-1]
    if ext not in SUPPORTED_EXTENSIONS:
        raise ParseError(
            f"Unsupported file type '{ext}'. Use one of: "
            + ", ".join(sorted(SUPPORTED_EXTENSIONS))
        )

    if ext == ".csv":
        text = decode_bytes(raw)
        rows = _parse_csv(text)
    else:  # .xlsx
        rows = _parse_xlsx(raw)

    if not rows:
        raise ParseError("The file contains no parseable rows (header required).")

    if len(rows) > ROW_CAP:
        raise ParseError(f"The file exceeds the {ROW_CAP}-row limit.")

    return rows
