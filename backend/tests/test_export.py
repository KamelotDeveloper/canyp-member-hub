"""Tests for the generic export endpoints (CSV + XLSX)."""

import io
from datetime import date

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from backend.models.pago import Pago
from backend.models.socio import Socio
from backend.services.importer.export_config import get_export_headers
from backend.services.importer.export_service import export_to_csv, export_to_xlsx

EXPORTABLE = ["socios", "pagos", "parcelas", "membresias", "aranceles", "notificaciones"]


def _seed_socios(db: Session) -> list[Socio]:
    """Insert two socios and return them."""
    s1 = Socio(
        id="s1", nombre="Ana Pérez", dni="30111111", telefono="11 5555 0101",
        email="ana@example.com", direccion="Calle 1", fechaAlta=date(2024, 1, 2), activo=True,
    )
    s2 = Socio(
        id="s2", nombre="Juan Gómez", dni="30222222", telefono="", email="",
        direccion="", fechaAlta=date(2024, 2, 3), activo=False,
    )
    db.add_all([s1, s2])
    db.commit()
    return [s1, s2]


class TestExportService:
    def test_csv_has_bom_headers_and_all_rows(self, test_db: Session):
        _seed_socios(test_db)
        csv_text = export_to_csv(test_db, "socios")
        assert csv_text.startswith("\ufeff")
        headers = [label for label, _ in get_export_headers("socios")]
        lines = csv_text.splitlines()
        assert headers == ["Nombre y Apellido", "DNI", "Teléfono", "Email", "Dirección", "Fecha de Alta", "Activo"]
        assert lines[1].startswith("Ana Pérez,30111111")
        assert lines[2].startswith("Juan Gómez,30222222")
        assert lines[1].endswith("Sí")
        assert lines[2].endswith("No")

    def test_xlsx_has_headers_and_rows(self, test_db: Session):
        _seed_socios(test_db)
        content = export_to_xlsx(test_db, "socios")
        wb = load_workbook(io.BytesIO(content))
        ws = wb.active
        assert ws["A1"].value == "Nombre y Apellido"
        assert ws["B1"].value == "DNI"
        assert ws["A2"].value == "Ana Pérez"
        assert ws["G2"].value == "Sí"

    def test_pagos_resolve_socio_nombre(self, test_db: Session):
        _seed_socios(test_db)
        pago = Pago(
            id="p1", numero="0001", socioId="s1", fecha=date(2024, 3, 4),
            medio="efectivo", total=100.0, nota="Cuota marzo",
        )
        test_db.add(pago)
        test_db.commit()
        csv_text = export_to_csv(test_db, "pagos")
        # The "Socio" column shows the readable name, not the raw id.
        assert "Ana Pérez" in csv_text
        assert "s1," not in csv_text
        headers = [label for label, _ in get_export_headers("pagos")]
        assert "Socio" in headers


class TestExportEndpoints:
    def test_unknown_resource_404(self, test_client):
        res = test_client.get("/api/export/inexistente")
        assert res.status_code == 404

    def test_invalid_format_422(self, test_client):
        res = test_client.get("/api/export/socios?format=pdf")
        assert res.status_code == 422

    def test_default_csv_export(self, test_client, test_db: Session):
        _seed_socios(test_db)
        res = test_client.get("/api/export/socios")
        assert res.status_code == 200
        assert res.headers["content-type"].startswith("text/csv")
        assert "attachment" in res.headers["content-disposition"]
        assert res.content.startswith(b"\xef\xbb\xbf")

    def test_xlsx_export_magic_bytes(self, test_client, test_db: Session):
        _seed_socios(test_db)
        res = test_client.get("/api/export/socios?format=xlsx")
        assert res.status_code == 200
        assert res.content[:2] == b"PK"
        assert "spreadsheetml" in res.headers["content-type"]
        wb = load_workbook(io.BytesIO(res.content))
        assert wb.active["A1"].value == "Nombre y Apellido"

    def test_all_resources_export_200(self, test_client, test_db: Session):
        _seed_socios(test_db)
        for resource in EXPORTABLE:
            res = test_client.get(f"/api/export/{resource}")
            assert res.status_code == 200, f"{resource} export failed: {res.text}"
            assert len(res.content) > 0

    def test_export_is_complete_not_paginated(self, test_client, test_db: Session):
        # Insert 3 socios; export must contain ALL of them regardless of any
        # pagination on the GET list endpoint.
        for i in range(3):
            test_db.add(
                Socio(
                    id=f"s{i}", nombre=f"Persona {i}", dni=f"30{i}000000",
                    fechaAlta=date(2024, 1, 1),
                )
            )
        test_db.commit()
        res = test_client.get("/api/export/socios")
        text = res.content.decode("utf-8-sig")
        for i in range(3):
            assert f"Persona {i}" in text